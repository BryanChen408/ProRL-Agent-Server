#!/usr/bin/env python3
"""一页 .pptx 生成器 —— 系统性能负载汇报(16:9,真 pptx,可编辑)。

内容定稿(用户口径):无 banner;KV/显存/aicore 用文字+表格,不画图;瓶颈排序不单列;
两个饼图(单 session 时间拆解 + 一轮累计)必须有;长度瓶颈三张直方图必须有;
引擎吞吐/时延表格必须有。
图由 matplotlib 生成 PNG 插入(本机无中文字体 → 图内一律英文标签);
文本框内中文由 PowerPoint 端字体渲染(Microsoft YaHei)。
数据计算复用 perf_report_html,与 report.html 数字严格一致。
用法: python3 gen_perf_pptx.py [--run <dir>] [--window-sessions N]
产物: <run_dir>/perf_report_html/perf_slide.pptx
"""
from __future__ import annotations
import argparse, glob, json, os, statistics as st, time
from collections import Counter, defaultdict
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from perf_report_html import _load, _q, _stats, _cv, _vt2, _bucket, _d

C = ["#2563eb", "#0ea5e9", "#f59e0b", "#10b981", "#ef4444", "#8b5cf6", "#f97316", "#94a3b8"]
NAVY = RGBColor(0x1E, 0x3A, 0x8A); INK = RGBColor(0x1F, 0x29, 0x37); MUT = RGBColor(0x64, 0x74, 0x8B)
ACC = RGBColor(0x25, 0x63, 0xEB); WHITE = RGBColor(0xFF, 0xFF, 0xFF); LINE = RGBColor(0xE5, 0xEA, 0xF2)
WARN = RGBColor(0xD9, 0x77, 0x06); BAD = RGBColor(0xDC, 0x26, 0x26)

plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 8.5,
                     "axes.edgecolor": "#dbe2ec", "axes.linewidth": 0.8})

# ---------- 图(matplotlib,英文标签) ----------
def fig_donut_session(path, sess_min, inf_min, pf_ratio, verify, tool, wait):
    """单 session 时间构成:双环 donut(外环标全部段名,内环仅标 Inference)。"""
    inner = [inf_min, verify, tool, wait]
    outer = [inf_min*pf_ratio, inf_min*(1-pf_ratio), verify, tool, wait]
    inner_lab = ["Inference", "", "", ""]
    outer_lab = ["Prefill", "Decode", "Verify", "Tool", "Wait"]
    inner_col = [C[0], C[2], C[3], C[7]]
    outer_col = [C[1], "#93c5fd", "#fbbf24", "#34d399", "#cbd5e1"]
    fig, ax = plt.subplots(figsize=(3.4, 3.4), dpi=220)
    ax.pie(outer, radius=1.0, colors=outer_col, labels=outer_lab, labeldistance=1.08,
           autopct=lambda p: f"{p:.0f}%" if p > 4 else "", pctdistance=0.85,
           wedgeprops=dict(width=0.30, edgecolor="w", linewidth=1.5),
           textprops=dict(fontsize=8.5, color="#334155"))
    ax.pie(inner, radius=0.68, colors=inner_col, labels=inner_lab, labeldistance=0.53,
           wedgeprops=dict(width=0.30, edgecolor="w", linewidth=1.5),
           textprops=dict(fontsize=8.5, color="white", weight="bold"))
    ax.text(0, 0.02, f"{sess_min:.0f} min", ha="center", va="center", fontsize=12.5, weight="bold", color="#0f172a")
    ax.text(0, -0.17, "per session", ha="center", va="center", fontsize=7.5, color="#64748b")
    ax.set(xlim=(-1.55, 1.55), ylim=(-1.42, 1.42))
    fig.subplots_adjust(left=0.01, right=0.99, top=0.99, bottom=0.01)
    fig.savefig(path, transparent=True); plt.close(fig)

def fig_donut_round(path, n, inf, verify, tool, wait, wall_h, conc):
    """一轮累计各段耗时 donut。"""
    vals = [inf, verify, tool, wait]
    labs = ["Inference", "Verify", "Tool", "Wait"]
    cols = [C[0], C[2], C[3], C[7]]
    fig, ax = plt.subplots(figsize=(3.4, 3.4), dpi=220)
    ax.pie(vals, radius=1.0, colors=cols, labels=labs, labeldistance=1.08,
           autopct=lambda p: f"{p:.0f}%" if p > 4 else "", pctdistance=0.83,
           wedgeprops=dict(width=0.34, edgecolor="w", linewidth=1.5),
           textprops=dict(fontsize=8.5, color="#334155"))
    ax.text(0, 0.06, f"{wall_h:.1f} h", ha="center", va="center", fontsize=13, weight="bold", color="#0f172a")
    ax.text(0, -0.12, f"{n} sessions", ha="center", va="center", fontsize=7.5, color="#64748b")
    ax.text(0, -0.26, f"conc. {conc:.0f}", ha="center", va="center", fontsize=7.5, color="#64748b")
    ax.set(xlim=(-1.55, 1.55), ylim=(-1.42, 1.42))
    fig.subplots_adjust(left=0.01, right=0.99, top=0.99, bottom=0.01)
    fig.savefig(path, transparent=True); plt.close(fig)

def fig_hist(path, vals, xlabel, vline=None, vline_lab=""):
    xs = sorted(v for v in vals if isinstance(v, (int, float)))
    fig, ax = plt.subplots(figsize=(4.3, 1.72), dpi=220)
    ax.hist(xs, bins=50, color=C[0], alpha=0.85, edgecolor="white", linewidth=0.3)
    if vline is not None:
        ax.axvline(vline, color="#ef4444", ls="--", lw=1)
        ax.text(vline, ax.get_ylim()[1]*0.92, f" {vline_lab}", color="#ef4444", fontsize=7.5, va="top")
    ax.set_xlabel(xlabel, fontsize=8); ax.set_ylabel("count", fontsize=8)
    ax.tick_params(labelsize=7.5); ax.grid(axis="y", color="#edf1f6", lw=0.7)
    ax.set_axisbelow(True)
    for s in ("top", "right"): ax.spines[s].set_visible(False)
    fig.subplots_adjust(left=0.1, right=0.98, top=0.95, bottom=0.24)
    fig.savefig(path, transparent=True); plt.close(fig)

# ---------- pptx 帮助函数 ----------
def _font(run, size, bold=False, color=INK, name="Microsoft YaHei"):
    f = run.font; f.size = Pt(size); f.bold = bold; f.color.rgb = color; f.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {}); rPr.append(ea)
    ea.set("typeface", name)

def txbox(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf

def para(tf, text, size, bold=False, color=INK, first=False, align=PP_ALIGN.LEFT, space_after=2):
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.alignment = align; p.space_after = Pt(space_after)
    r = p.add_run(); r.text = text; _font(r, size, bold, color)
    return p

def panel(slide, x, y, w, h):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = 0.045
    sh.fill.solid(); sh.fill.fore_color.rgb = WHITE
    sh.line.color.rgb = LINE; sh.line.width = Pt(1)
    sh.shadow.inherit = False
    return sh

def kpi(slide, x, y, w, h, val, lab):
    panel(slide, x, y, w, h)
    tf = txbox(slide, x, y + 0.07, w, h - 0.1)
    para(tf, val, 16, True, ACC, True, PP_ALIGN.CENTER, 0)
    para(tf, lab, 8.5, False, MUT, align=PP_ALIGN.CENTER, space_after=0)

# ---------- 主流程 ----------
def build(run_dir, metrics_dir, window_n):
    out = os.path.join(run_dir, "perf_report_html"); os.makedirs(out, exist_ok=True)
    der = os.path.join(run_dir, "telemetry_derived")
    assets = os.path.join(out, "pptx_assets"); os.makedirs(assets, exist_ok=True)

    # ===== 数据加载(与 report/slide 相同) =====
    ses = []
    for f in glob.glob(f"{run_dir}/rollout_results/**/ses_*.json", recursive=True):
        try: s = json.load(open(f))
        except Exception: continue
        md=(s.get("trajectory",{}) or {}).get("metadata",{}) or {}; tm=md.get("task_metadata") or {}
        t=s.get("timing",{}) or {}; rm=t.get("run_ms") or 0; end=os.path.getmtime(f)
        ses.append(dict(op=tm.get("op_name"), status=s.get("status"), run_ms=rm,
                        start=end-rm/1000, end=end))
    if not ses: raise SystemExit("无完成 session")
    ses.sort(key=lambda s: s["end"])
    if window_n: ses = ses[-window_n:]
    W0, W1 = min(s["start"] for s in ses), max(s["end"] for s in ses)
    inwin = lambda rows: [r for r in rows if W0 <= r.get("recorded_at_unix",0) <= W1]
    t1=inwin(_load(f"{metrics_dir}/npu_state/npu_card.jsonl"))
    t2={e:_vt2(inwin(_load(f"{metrics_dir}/vllm_state/{e}.jsonl"))) for e in ("infer-0","infer-1","infer-2")}
    t4=[r for f in glob.glob(f"{metrics_dir}/engine-*.jsonl") for r in _load(f)]
    t5=_load(f"{der}/rollout_span.jsonl"); t6=_load(f"{der}/verify_job.jsonl")

    # ===== 计算(与 report 相同口径) =====
    stc=Counter(s["status"] for s in ses)
    rm_all=[s["run_ms"] for s in ses if s.get("run_ms")]
    allkv=[r.get("vllm:kv_cache_usage_perc",0)*100 for rows in t2.values() for r in rows]
    allwt=[r.get("vllm:num_requests_waiting",0) for rows in t2.values() for r in rows]
    last=ses[-64:] if len(ses)>=64 else ses
    wall_h=(max(s["end"] for s in last)-min(s["start"] for s in last))/3600
    sum_h=sum(s["run_ms"] for s in last if s.get("run_ms"))/1000/3600
    conc=sum_h/wall_h if wall_h else 0
    eng_rows=[]; tps={}
    for e, rows in t2.items():
        if len(rows)<2: continue
        dt=rows[-1]["recorded_at_unix"]-rows[0]["recorded_at_unix"]
        tp=_d(rows,"vllm:generation_tokens_total")/dt if dt else 0; tps[e]=tp
        ttft=_d(rows,"vllm:time_to_first_token_seconds_sum")/max(1,_d(rows,"vllm:time_to_first_token_seconds_count"))
        itl=_d(rows,"vllm:inter_token_latency_seconds_sum")/max(1,_d(rows,"vllm:inter_token_latency_seconds_count"))
        kv=[r.get("vllm:kv_cache_usage_perc",0)*100 for r in rows]
        wt=[r.get("vllm:num_requests_waiting",0) for r in rows]
        ph=_d(rows,"vllm:prefix_cache_hits_total"); pq=_d(rows,"vllm:prefix_cache_queries_total")
        eng_rows.append((e,tp,ttft,itl*1000,st.mean(kv),max(kv),st.mean(wt),100*ph/pq if pq else 0))
    ttft_tot=e2e_tot=0
    for rows in t2.values():
        if len(rows)>=2:
            ttft_tot+=_d(rows,"vllm:time_to_first_token_seconds_sum")
            e2e_tot+=_d(rows,"vllm:e2e_request_latency_seconds_sum")
    pf_ratio=ttft_tot/e2e_tot if e2e_tot else 0.15
    means={}
    for k,cn in [("infer_frac","推理"),("verify_frac","验证"),("tool_frac","工具"),("wait_cpu_frac","等待/CPU")]:
        vals=[s.get(k) for s in t5 if isinstance(s.get(k),(int,float))]; means[cn]=st.mean(vals) if vals else 0
    sess_min=st.mean(rm_all)/60000 if rm_all else 0
    inf_min=means["推理"]*sess_min
    inf=[r.get("aicore_util_pct") for r in t1 if r["pool"]=="inference" and r.get("aicore_util_pct") is not None]
    ver=[r.get("aicore_util_pct") for r in t1 if r["pool"]=="verify" and r.get("aicore_util_pct") is not None]
    hbm=[r.get("hbm_used_mb") for r in t1 if r["pool"]=="inference" and r.get("hbm_used_mb")]
    pt=[r.get("num_prompt_tokens") for r in t4]; gt=[r.get("num_generation_tokens") for r in t4]
    ctx=[(r.get("num_prompt_tokens") or 0)+(r.get("num_generation_tokens") or 0) for r in t4 if r.get("num_prompt_tokens")]
    ch=[r.get("prefix_cache_hit_pct") for r in t4]
    over=sum(1 for c in ctx if c>240000)
    noninf=100-st.mean([s.get('infer_frac',0) for s in t5])*100 if t5 else 0
    et=Counter(v.get("error_type") for v in t6) if t6 else Counter()

    # ===== 图 PNG =====
    fig_donut_session(f"{assets}/donut_session.png", sess_min, inf_min, pf_ratio,
                      means["验证"]*sess_min, means["工具"]*sess_min, means["等待/CPU"]*sess_min)
    n=len(last)
    fig_donut_round(f"{assets}/donut_round.png", n, means["推理"]*sess_min*n, means["验证"]*sess_min*n,
                    means["工具"]*sess_min*n, means["等待/CPU"]*sess_min*n, wall_h, conc)
    fig_hist(f"{assets}/hist_prompt.png", pt, "prompt tokens")
    fig_hist(f"{assets}/hist_decode.png", gt, "generation tokens")
    fig_hist(f"{assets}/hist_ctx.png", ctx, "total context tokens", vline=262144, vline_lab="262144 limit")

    # ===== pptx =====
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor(0xF3, 0xF5, 0xF9)

    # 标题(无 banner,纯文字)
    tf = txbox(slide, 0.4, 0.14, 12.5, 0.42)
    para(tf, "ProRL 推理侧系统负载分析", 21, True, NAVY, True, space_after=0)
    tf = txbox(slide, 0.4, 0.56, 12.5, 0.26)
    para(tf, f"run {os.path.basename(run_dir.rstrip('/'))} · {len(ses)} session · "
             f"{time.strftime('%m-%d %H:%M',time.localtime(W0))}→{time.strftime('%H:%M',time.localtime(W1))} · "
             f"拓扑:16 NPU = 12 推理(3 引擎×TP4)+ 4 验证",
         9.5, False, MUT, True, space_after=0)

    # KPI 条
    kpis = [(f"{stc.get('COMPLETED',0)/len(ses)*100:.0f}%", "goodput"),
            (f"{(_q(rm_all,.5) or 0)/60000:.0f} min", "单 session p50"),
            (f"{st.mean(allkv):.0f}%", "KV cache 均值"),
            (f"{wall_h:.2f} h", "一轮墙钟(64 session)"),
            (f"{conc:.1f}", "并发度(≈12 推理卡)"),
            (f"{len(set(s['op'] for s in ses))}", "覆盖算子")]
    kw, kg = 2.03, 0.073
    for i, (v, l) in enumerate(kpis):
        kpi(slide, 0.4 + i*(kw+kg), 0.92, kw, 0.62, v, l)

    # ===== 左:引擎性能 + 饱和度文字 =====
    panel(slide, 0.4, 1.72, 4.72, 2.86)
    tf = txbox(slide, 0.58, 1.84, 4.4, 0.28)
    para(tf, "引擎性能(吞吐 / 时延)", 12.5, True, INK, True, space_after=0)
    rows, cols = len(eng_rows)+1, 8
    tbl_shape = slide.shapes.add_table(rows, cols, Inches(0.56), Inches(2.16), Inches(4.4), Inches(0.36*rows))
    tbl = tbl_shape.table
    heads = ["引擎", "吞吐\ntok/s", "TTFT\ns", "ITL\nms", "KV均\n%", "KV峰\n%", "wait\n均", "prefix\n%"]
    for j, htx in enumerate(heads):
        c = tbl.cell(0, j); c.text = ""
        p = c.text_frame.paragraphs[0]; r = p.add_run(); r.text = htx
        _font(r, 8, True, MUT); p.alignment = PP_ALIGN.CENTER
    for i, (e, tp, ttft, itl, kva, kvp, wt, pf) in enumerate(eng_rows, 1):
        vals = [e, f"{tp:.0f}", f"{ttft:.2f}", f"{itl:.1f}", f"{kva:.0f}", f"{kvp:.0f}", f"{wt:.2f}", f"{pf:.0f}"]
        for j, v in enumerate(vals):
            c = tbl.cell(i, j); c.text = ""
            p = c.text_frame.paragraphs[0]; r = p.add_run(); r.text = v
            _font(r, 9, j == 1, INK if j else MUT); p.alignment = PP_ALIGN.CENTER
    cv = _cv(list(tps.values()))
    tf = txbox(slide, 0.58, 3.72, 4.4, 0.8)
    para(tf, f"· 引擎均衡 CV={cv} — {'不均,LB/DP 分发不平' if cv>0.1 else '三引擎均衡'}", 9.5, False, INK, True)
    para(tf, f"· KV 均值 {st.mean(allkv):.0f}%、waiting≈{st.mean(allwt):.2f} → 引擎未饱和,rollout 并发可大幅提高", 9.5)
    para(tf, f"· aicore:推理池均 {st.mean(inf or [0]):.0f}% / 验证池均 {st.mean(ver or [0]):.0f}%(间歇打满);"
             f"HBM 占满约 {(_stats(hbm)['mean'] if hbm else 0)/65536*100:.0f}%(权重+激活占大头,KV 有空间)", 9.5)

    # ===== 中:单 session 时间构成 =====
    panel(slide, 5.24, 1.72, 3.95, 2.86)
    tf = txbox(slide, 5.42, 1.84, 3.6, 0.28)
    para(tf, "单 session 时间构成(均值)", 12.5, True, INK, True, space_after=0)
    slide.shapes.add_picture(f"{assets}/donut_session.png", Inches(6.32), Inches(2.1), Inches(1.8), Inches(1.8))
    tf = txbox(slide, 5.42, 3.94, 3.6, 0.6)
    para(tf, f"prefill:decode ≈ {pf_ratio*100:.0f}:{100-pf_ratio*100:.0f}(引擎侧近似)", 9, False, INK, True)
    para(tf, f"推理外开销 ≈ {100-means['推理']*100:.0f}%(≈{sess_min-inf_min:.0f} min/session)", 9)

    # ===== 右:一轮累计 =====
    panel(slide, 9.31, 1.72, 3.62, 2.86)
    tf = txbox(slide, 9.49, 1.84, 3.3, 0.28)
    para(tf, f"一轮({n} session)累计耗时", 12.5, True, INK, True, space_after=0)
    slide.shapes.add_picture(f"{assets}/donut_round.png", Inches(10.22), Inches(2.1), Inches(1.8), Inches(1.8))
    tf = txbox(slide, 9.49, 3.94, 3.3, 0.6)
    para(tf, f"墙钟 {wall_h:.2f}h · 累加 {sum_h:.0f}h · 并发度 {conc:.1f}", 9, False, INK, True)
    para(tf, "缩短一轮 = 缩短单 session 或提并发", 9)

    # ===== 下:长度瓶颈三图 =====
    panel(slide, 0.4, 4.72, 12.53, 2.14)
    tf = txbox(slide, 0.58, 4.82, 8.0, 0.26)
    para(tf, "长度瓶颈(逐请求分布)", 12.5, True, INK, True, space_after=0)
    slide.shapes.add_picture(f"{assets}/hist_prompt.png", Inches(0.62), Inches(5.12), Inches(4.0), Inches(1.6))
    slide.shapes.add_picture(f"{assets}/hist_decode.png", Inches(4.72), Inches(5.12), Inches(4.0), Inches(1.6))
    slide.shapes.add_picture(f"{assets}/hist_ctx.png", Inches(8.82), Inches(5.12), Inches(4.0), Inches(1.6))
    tf = txbox(slide, 0.58, 6.6, 12.2, 0.24)
    para(tf, f"prompt p50 {(_q(pt,.5) or 0)/1000:.0f}k tok(长 prompt 是主体)· decode p50 {_q(gt,.5) or 0:.0f} tok(生成压力小)· "
             f"context>240k 占 {100*over/max(1,len(ctx)):.1f}% · prefix 命中 p50 {_q(ch,.5) or 0:.0f}%(prefill 大量走缓存)",
         9, False, MUT, True, space_after=0)

    # ===== 底:结论一行 =====
    top_err = f"{et.most_common(1)[0][0]}({100*et.most_common(1)[0][1]/len(t6):.0f}%)" if et else "—"
    tf = txbox(slide, 0.4, 6.98, 12.53, 0.44)
    para(tf, f"结论:① 通过率瓶颈 = {top_err}(第一杠杆)② 单 session 长(p50 {(_q(rm_all,.5) or 0)/60000:.0f}min,主耗在 agent 多轮)"
             f"③ 推理算力过剩(KV {st.mean(allkv):.0f}%,可提并发)④ agent 侧开销 {noninf:.0f}% ⑤ 验证池闲置(4 卡均 {st.mean(ver or [0]):.0f}%)",
         9.5, False, INK, True, space_after=0)
    para(tf, "口径:卡负载 5s · 引擎 metrics 差分 · 逐请求日志 · session 时间拆解(source_data/ 可复核)",
         7.5, False, MUT, space_after=0)

    rep = f"{out}/perf_slide.pptx"
    prs.save(rep)
    return rep

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--run")
    ap.add_argument("--runs-root", default=os.environ.get("POLAR_RUNS_ROOT",
        "/mnt/share/c00937190/src/cannbot_debug/ProRL-Agent-Server/output/ascend_operator/runs"))
    ap.add_argument("--metrics-dir", default=os.environ.get("POLAR_ENGINE_METRICS_DIR","/mnt/share/polar_engine_metrics"))
    ap.add_argument("--window-sessions", type=int, default=0)
    a=ap.parse_args()
    run=a.run or max(glob.glob(f"{a.runs_root}/*/"), key=os.path.getmtime)
    print(f"PPTX: {build(run, a.metrics_dir, a.window_sessions)}")

if __name__=="__main__":
    main()
